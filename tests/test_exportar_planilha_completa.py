import io
import pandas as pd
import openpyxl
from pages import painel


def test_exportar_planilha_completa_cria_abas_e_formatos():
    df = pd.DataFrame(
        {
            "Tipo Nota": ["Entrada", "Saída"],
            "Valor Total": [1000.5, 2000.75],
            "ICMS Alíquota": [0.1, 0.2],
        }
    )

    data = painel._exportar_planilha_completa(df)
    wb = openpyxl.load_workbook(io.BytesIO(data))

    assert set(wb.sheetnames) == {"entradas", "saídas"}

    ws = wb["entradas"]
    assert ws["B2"].value == 1000.5
    assert ws["B2"].number_format == "R$ #,##0.00"
    assert ws["C2"].value == 0.1
    assert ws["C2"].number_format == "0.00%"
